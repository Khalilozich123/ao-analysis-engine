"""Parse an uploaded .xlsx into normalized Opportunity records.

Handles the real procurement export: ~60 columns, accented French headers, and dates stored as
Excel serial numbers. Unknown columns are ignored; rows missing required values are skipped;
a workbook missing a required COLUMN is rejected (raises ValueError) so there's no partial run.
"""

import io
import re
import unicodedata
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from app.schemas.opportunity import Opportunity


def _norm_header(h) -> str:
    """Lowercase, strip accents, collapse non-alphanumerics to single spaces."""
    s = unicodedata.normalize("NFKD", str(h or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


# Normalized header -> target field. `objet` is special (fills both title and description).
HEADER_MAP = {
    "n ordre": "id", "no ordre": "id", "numero d ordre": "id", "numero dordre": "id", "numero": "id",
    "reference": "reference", "ref": "reference",
    "objet": "objet",
    "organisme": "client",
    "ville": "city",
    "budget": "budget",
    "caution": "deposit",
    "date limite": "deadline", "date limite de soumission": "deadline",
}

REQUIRED_TARGETS = {"id", "reference", "objet"}  # must exist as columns
_TITLE_MAX = 120


def _coerce_date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):  # Excel serial number
        try:
            d = from_excel(v)
            return d.date() if isinstance(d, datetime) else d
        except Exception:
            return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            continue
    return None


def _clean(v):
    return str(v).strip() if v not in (None, "") else None


def parse_opportunities(data: bytes) -> list[Opportunity]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)

    try:
        header = next(rows)
    except StopIteration:
        raise ValueError("Le fichier est vide.")

    # Map each column index to a target field (ignore unknown columns).
    col_to_target: dict[int, str] = {}
    for idx, cell in enumerate(header):
        target = HEADER_MAP.get(_norm_header(cell))
        if target and target not in col_to_target.values():
            col_to_target[idx] = target

    found = set(col_to_target.values())
    missing = REQUIRED_TARGETS - found
    if missing:
        pretty = {"id": "Numéro d'ordre", "reference": "Référence", "objet": "Objet"}
        raise ValueError(
            "Colonnes requises manquantes : " + ", ".join(pretty[m] for m in sorted(missing))
        )

    opportunities: list[Opportunity] = []
    seen: set[str] = set()
    for row in rows:
        fields: dict = {}
        for idx, target in col_to_target.items():
            val = row[idx] if idx < len(row) else None
            if target == "deadline":
                fields["deadline"] = _coerce_date(val)
            elif target == "objet":
                objet = _clean(val)
                if objet:
                    fields["description"] = objet
                    fields["title"] = objet[:_TITLE_MAX]
            else:
                fields[target] = _clean(val)

        opp_id = fields.get("id")
        # Skip rows missing required values, and de-duplicate within the file on id.
        if not opp_id or not fields.get("reference") or not fields.get("description"):
            continue
        if opp_id in seen:
            continue
        seen.add(opp_id)

        try:
            opportunities.append(Opportunity(**fields))
        except Exception:
            continue  # a single bad row never fails the whole upload

    return opportunities
