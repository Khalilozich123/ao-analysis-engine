"""Build an .xlsx export of analyzed results (matches the PowerApps export columns)."""

import io

from openpyxl import Workbook
from openpyxl.styles import Font

COLUMNS = [
    "Rang", "N° ordre", "Référence", "Objet", "Organisme", "Ville",
    "Budget", "Caution", "Score", "Verdict", "Action", "Justification",
]
WIDTHS = [6, 12, 16, 48, 34, 14, 14, 12, 8, 16, 16, 70]


def build_xlsx(rows) -> bytes:
    """rows: Result ORM objects, already ordered (rank = position)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Analyse"

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    for rank, r in enumerate(rows, 1):
        ws.append([
            rank, r.opportunity_id, r.reference, r.title, r.client, r.city,
            r.budget, r.deposit, r.score, r.verdict, r.action, r.rationale,
        ])

    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
